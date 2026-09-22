#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
from dotenv import load_dotenv
load_dotenv()

import asyncio
import glob
import time
import json
import pickle
import hashlib
import sqlite3
import subprocess
import shutil
import re
import uuid
import random
import socket
from datetime import datetime, timedelta, timezone
from contextlib import contextmanager
from collections import defaultdict
from typing import List, Dict, Optional, Tuple, Any

import numpy as np
import pandas as pd
import torch
import cv2
import easyocr
from openai import OpenAI
from sentence_transformers import SentenceTransformer
import faiss
from pypdf import PdfReader
from docx import Document

from fastapi import FastAPI, HTTPException, Depends, Request, UploadFile, File
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from pydantic import BaseModel, Field
from jose import JWTError, jwt
from passlib.context import CryptContext
import uvicorn
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded

app = FastAPI(title="RAG Экстрадиционный Анализатор")
limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

os.environ["VLLM_USE_FLASHINFER_SAMPLER"] = "0"

# ====================== КОНФИГУРАЦИЯ ======================
SECRET_KEY = os.getenv("SECRET_KEY")
assert SECRET_KEY, "SECRET_KEY not loaded"
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60

VLLM_HOST = "0.0.0.0"
VLLM_PORT = 8001
VLLM_API_KEY = os.getenv("VLLM_API_KEY", "rag_system_key_2026")
SERVED_MODEL_NAME = "YandexGPT-5-Lite-8B-instruct"
BASE_MODEL = "yandex/YandexGPT-5-Lite-8B-instruct"

RAG_BASE_FOLDER = "./RAG_BASE"
UPLOAD_ROOT = "./uploads"
CHUNK_SIZE = 800
CHUNK_OVERLAP = 100
TOP_K = 3

RAG_INDEX_CACHE = "./rag_faiss.index"
RAG_META_CACHE = "./rag_meta.pkl"

# ====================== SYSTEM PROMPT ======================
SYSTEM_PROMPT = (
    "Ты — эксперт по экстрадиционному праву. Ведёшь живой диалог с юристом.\n\n"
    "ПРАВИЛА ДИАЛОГА:\n"
    "0. Если в диалоге ниже есть предыдущие реплики — используй их в первую очередь.\n"
    "0.1. На вопросы «какой был мой предыдущий вопрос», «о чём мы говорили», "
    "«что я спрашивал» — отвечай ИЗ ИСТОРИИ диалога, дословно цитируя реплики пользователя.\n"
    "0.2. Местоимения («он», «она», «это», «эта», «тот», «там», «ещё») "
    "раскрывай из последних 3–5 реплик диалога.\n"
    "0.3. Если формулировка грубая или непонятно, о ком речь — вежливо уточни, "
    "о ком именно идёт вопрос. Не отказывайся отвечать.\n"
    "0.4. НЕ начинай каждый ответ с «В представленных документах...». "
    "Отвечай как человек: коротко, по делу, если вопрос уточняющий.\n"
    "0.5. НЕ дублируй дословно предыдущий ответ. Отвечай на новую реплику.\n"
    "1. Если ответ есть в истории — отвечай из неё, не говори «нет информации».\n"
    "2. Ссылайся на статьи, но не превращай каждый ответ в отчёт.\n"
    "3. Если данных правда нет — скажи одной фразой и предложи, что проверить.\n"
    "4. Можешь формировать аналитические справки, заключения, "
    "перечни замечаний, сравнительные таблицы — по запросу.\n"
    "5. Пиши на русском, официально-деловым или консультативным стилем "
    "в зависимости от вопроса."
)

# ====================== ИНИЦИАЛИЗАЦИЯ ======================
os.makedirs(RAG_BASE_FOLDER, exist_ok=True)
os.makedirs(UPLOAD_ROOT, exist_ok=True)

vllm_process = None
vector_store = None
client = None
ocr_reader = None

DATABASE_FILE = "rag_system.db"

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
security = HTTPBearer()

# ====================== БАЗА ДАННЫХ ======================
@contextmanager
def get_db():
    conn = sqlite3.connect(DATABASE_FILE)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()

def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS users (
                username TEXT PRIMARY KEY,
                hashed_password TEXT NOT NULL,
                full_name TEXT,
                organization TEXT,
                position TEXT,
                registered_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS upload_sessions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL,
                session_id TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(username) REFERENCES users(username)
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS uploaded_docs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                session_id TEXT NOT NULL,
                filename TEXT NOT NULL,
                file_path TEXT NOT NULL,
                doc_type TEXT,
                extracted_text TEXT,
                entities TEXT,
                uploaded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS reports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL,
                session_id TEXT NOT NULL,
                report_text TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(username) REFERENCES users(username)
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_session ON uploaded_docs(session_id)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_reports_user ON reports(username)")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS chats (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                chat_id TEXT NOT NULL,
                username TEXT NOT NULL,
                session_id TEXT,
                title TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(username) REFERENCES users(username)
            )
        """)
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_chats_chat_id ON chats(chat_id)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_chats_user ON chats(username, updated_at DESC)")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS chat_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                chat_id TEXT NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(chat_id) REFERENCES chats(chat_id)
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_messages_chat ON chat_messages(chat_id, id ASC)")

def get_user(username: str) -> Optional[dict]:
    with get_db() as conn:
        row = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        return dict(row) if row else None

def create_user(username: str, hashed_password: str, full_name: str = None,
                organization: str = None, position: str = None):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO users (username, hashed_password, full_name, organization, position) "
            "VALUES (?, ?, ?, ?, ?)",
            (username, hashed_password, full_name, organization, position)
        )

# ====================== АУТЕНТИФИКАЦИЯ ======================
def verify_password(plain: str, hashed: str) -> bool:
    return pwd_context.verify(plain, hashed)

def get_password_hash(password: str) -> str:
    password_bytes = password.encode('utf-8')[:72]
    return pwd_context.hash(password_bytes.decode('utf-8', errors='ignore'))

def authenticate_user(username: str, password: str) -> Optional[dict]:
    user = get_user(username)
    if not user or not verify_password(password, user["hashed_password"]):
        return None
    return user

def create_access_token(data: dict, expires_delta: timedelta = None):
    to_encode = data.copy()
    expire = datetime.now(timezone.utc) + (expires_delta or timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    token = credentials.credentials
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise HTTPException(status_code=401, detail="Invalid token")
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid token")
    user = get_user(username)
    if user is None:
        raise HTTPException(status_code=401, detail="User not found")
    return username

# ====================== vLLM СЕРВЕР ======================
def start_vllm_server():
    global vllm_process
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    result = sock.connect_ex(('localhost', VLLM_PORT))
    sock.close()
    if result == 0:
        print(f"vLLM сервер уже запущен на порту {VLLM_PORT}")
        return

    print("Запуск vLLM сервера...")
    command = [
        "vllm", "serve", BASE_MODEL,
        "--host", VLLM_HOST,
        "--port", str(VLLM_PORT),
        "--api-key", VLLM_API_KEY,
        "--served-model-name", SERVED_MODEL_NAME,
        "--max-model-len", "32768",
        "--max-num-seqs", "4",
        "--gpu-memory-utilization", "0.90",
        "--dtype", "bfloat16",
    ]
    log_file = open("vllm_rag.log", "w")
    vllm_process = subprocess.Popen(command, stdout=log_file, stderr=subprocess.STDOUT,
                                    text=True, start_new_session=True)
    while True:
        time.sleep(1)
        with open("vllm_rag.log", "r") as f:
            if "Application startup complete" in f.read():
                print("vLLM сервер готов")
                break

# ====================== RAG ВЕКТОРНОЕ ХРАНИЛИЩЕ ======================
class VectorStore:
    def __init__(self, model_name="Roflmax/bge-m3-russian-legal"):
        self.embedder = SentenceTransformer(model_name, device="cpu")
        self.index = None
        self.chunks = []
        self.metadata = []

    def add_documents(self, docs: List[Tuple[str, str]]):
        all_chunks, all_meta = [], []
        for filename, text in docs:
            for i, ch in enumerate(self._chunk_text(text)):
                all_chunks.append(ch)
                all_meta.append((filename, i))
        if not all_chunks:
            return
        print(f"Генерация эмбеддингов для {len(all_chunks)} чанков...")
        embeddings = self.embedder.encode(all_chunks, show_progress_bar=True)
        self.index = faiss.IndexFlatL2(embeddings.shape[1])
        self.index.add(embeddings.astype(np.float32))
        self.chunks = all_chunks
        self.metadata = all_meta
        print(f"Индекс создан. Всего векторов: {self.index.ntotal}")

    def _chunk_text(self, text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
        words = text.split()
        chunks, start = [], 0
        while start < len(words):
            end = min(start + chunk_size, len(words))
            chunks.append(" ".join(words[start:end]))
            if end == len(words):
                break
            start = end - overlap
        return chunks

    def save_cache(self, index_path=RAG_INDEX_CACHE, meta_path=RAG_META_CACHE, signature=None):
        if self.index is None:
            return
        faiss.write_index(self.index, index_path)
        with open(meta_path, "wb") as f:
            pickle.dump({"chunks": self.chunks, "metadata": self.metadata,
                         "signature": signature}, f)

    def load_cache(self, index_path=RAG_INDEX_CACHE, meta_path=RAG_META_CACHE):
        self.index = faiss.read_index(index_path)
        with open(meta_path, "rb") as f:
            data = pickle.load(f)
        self.chunks = data["chunks"]
        self.metadata = data["metadata"]
        return data.get("signature")

    def search(self, query, top_k=TOP_K):
        if self.index is None:
            return []
        q_emb = self.embedder.encode([query])
        distances, indices = self.index.search(q_emb.astype(np.float32), top_k)
        results = []
        for dist, idx in zip(distances[0], indices[0]):
            if idx != -1:
                results.append((self.chunks[idx], self.metadata[idx][0], dist))
        return results

def load_legal_documents(folder: str) -> List[Tuple[str, str]]:
    docs = []
    for file_path in glob.glob(os.path.join(folder, "*")):
        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext == ".pdf":
                reader = PdfReader(file_path)
                text = "\n".join([p.extract_text() or "" for p in reader.pages])
            elif ext == ".docx":
                doc = Document(file_path)
                text = "\n".join([p.text for p in doc.paragraphs])
            elif ext == ".txt":
                with open(file_path, "r", encoding="utf-8") as f:
                    text = f.read()
            else:
                continue
            if text.strip():
                docs.append((os.path.basename(file_path), text))
        except Exception as e:
            print(f"Ошибка загрузки {file_path}: {e}")
    return docs

def rag_base_signature(folder: str) -> str:
    items = []
    for file_path in sorted(glob.glob(os.path.join(folder, "*"))):
        if not os.path.isfile(file_path):
            continue
        ext = os.path.splitext(file_path)[1].lower()
        if ext not in (".pdf", ".docx", ".txt"):
            continue
        st = os.stat(file_path)
        items.append(f"{os.path.basename(file_path)}:{st.st_size}:{int(st.st_mtime)}")
    return hashlib.sha256("\n".join(items).encode("utf-8")).hexdigest()

def init_rag():
    global vector_store
    signature = rag_base_signature(RAG_BASE_FOLDER)
    vector_store = VectorStore()
    if os.path.exists(RAG_INDEX_CACHE) and os.path.exists(RAG_META_CACHE):
        try:
            cached_signature = vector_store.load_cache()
            if cached_signature == signature:
                print(f"RAG индекс загружен из кэша: {vector_store.index.ntotal}")
                return
            print("RAG_BASE изменился, пересобираю...")
        except Exception as e:
            print(f"Кэш не загрузился: {e}")
    legal_docs = load_legal_documents(RAG_BASE_FOLDER)
    if not legal_docs:
        print("ВНИМАНИЕ: RAG_BASE пуста.")
        return
    vector_store.add_documents(legal_docs)
    vector_store.save_cache(signature=signature)
    print("RAG индекс сохранён")

# ====================== OCR / ИЗВЛЕЧЕНИЕ ТЕКСТА ======================
def get_ocr_reader():
    global ocr_reader
    if ocr_reader is None:
        print("Загрузка EasyOCR...")
        ocr_reader = easyocr.Reader(['ru', 'en'], gpu=False)
    return ocr_reader

def extract_text_from_file(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            return "\n".join([p.extract_text() or "" for p in reader.pages])
        elif ext == ".docx":
            doc = Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs])
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
            reader = get_ocr_reader()
            img = cv2.imread(file_path)
            if img is None:
                return ""
            gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
            gray = cv2.resize(gray, None, fx=1.5, fy=1.5, interpolation=cv2.INTER_CUBIC)
            return "\n".join(reader.readtext(gray, detail=0, paragraph=True, width_ths=0.7)).strip()
        else:
            return ""
    except Exception as e:
        print(f"Ошибка извлечения из {file_path}: {e}")
        return ""

# ====================== КЛАССИФИКАЦИЯ И СУЩНОСТИ ======================
DOCUMENT_TYPES = [
    "запрос о выдаче", "постановление о привлечении в качестве обвиняемого",
    "постановление суда о заключении под стражу", "постановление об объявлении в розыск",
    "текст статей уголовного закона запрашивающего государства",
    "справка о гражданстве", "справка о статусе беженца",
    "протокол опроса лица", "справка о судимости",
    "справка об отсутствии уголовного преследования", "копия паспорта",
    "копия миграционной карты", "справка о регистрации",
    "заключение прокуратуры субъекта РФ", "постановление суда об избрании меры пресечения",
    "нотариально заверенный перевод запроса", "гарантии запрашивающего государства", "иное"
]

def classify_document(text: str, filename: str) -> str:
    text_lower = text.lower()[:2000]
    fn = filename.lower()
    if "запрос" in fn and ("выдач" in fn or "экстрадиц" in fn):
        return "запрос о выдаче"
    if "постановление о привлечении" in text_lower or "привлечении в качестве обвиняемого" in text_lower:
        return "постановление о привлечении в качестве обвиняемого"
    if "заключение под стражу" in text_lower:
        return "постановление суда о заключении под стражу"
    if "объявлен в розыск" in text_lower or "розыск" in text_lower:
        return "постановление об объявлении в розыск"
    if "справка" in fn and "гражданств" in fn:
        return "справка о гражданстве"
    if "протокол опроса" in text_lower or "опрос" in text_lower:
        return "протокол опроса лица"
    if "паспорт" in fn:
        return "копия паспорта"
    if "миграционная карта" in fn:
        return "копия миграционной карты"
    if "заключение прокуратуры" in text_lower:
        return "заключение прокуратуры субъекта РФ"
    if "мера пресечения" in text_lower or "заключение под стражу" in text_lower:
        return "постановление суда об избрании меры пресечения"
    if "гара" in fn or "гарант" in text_lower:
        return "гарантии запрашивающего государства"
    return "иное"

def extract_entities_from_document(text: str, doc_type: str, llm_client) -> Dict[str, Any]:
    prompt = f"""
Ты - система извлечения информации из экстрадиционных документов.
Извлеки поля: ФИО, Дата рождения, Место рождения, Гражданство, Паспортные данные,
Квалификация деяния, Запрашивающее государство, Запрашивающий орган, Дата запроса,
Номер запроса, Период совершения деяния, Место совершения деяния, Мера наказания,
Дата постановления о привлечении, Дата постановления о заключении под стражу,
Дата избрания меры пресечения, Срок меры пресечения, Отметка о согласии/несогласии,
Заявления лица.

Документ:
{text}

Ответ дай в формате JSON. Если поле не найдено — null.
"""
    try:
        response = llm_client.chat.completions.create(
            model=SERVED_MODEL_NAME,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.0, max_tokens=8192,
        )
        content = response.choices[0].message.content
        m = re.search(r'\{.*\}', content, re.DOTALL)
        return json.loads(m.group()) if m else {}
    except Exception as e:
        print(f"Ошибка извлечения сущностей: {e}")
        return {}

# ====================== КРОСС-ПРОВЕРКА ======================
def cross_check_entities(all_entities: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    discrepancies = []
    field_values = defaultdict(list)
    for i, ents in enumerate(all_entities):
        for key, value in ents.items():
            if value and value != "null" and value is not None:
                if isinstance(value, (dict, list)):
                    value_str = json.dumps(value, ensure_ascii=False, sort_keys=True)
                else:
                    value_str = str(value)
                field_values[key].append((value_str, i, value))
    for key, values in field_values.items():
        if len(set(v[0] for v in values)) > 1:
            discrepancies.append({
                "field": key,
                "values": [{"value": v[2], "doc_index": v[1]} for v in values]
            })
    return discrepancies

# ====================== ПРАВОВОЙ АНАЛИЗ ======================
def legal_analysis(extracted_data: Dict[str, Any], discrepancies: List[Dict], llm_client) -> Dict[str, Any]:
    context_info = f"""
Лицо: {extracted_data.get('ФИО', 'не указано')}
Гражданство: {extracted_data.get('Гражданство', 'не указано')}
Дата рождения: {extracted_data.get('Дата рождения', 'не указано')}
Квалификация: {extracted_data.get('Квалификация деяния', 'не указано')}
Период: {extracted_data.get('Период совершения деяния', 'не указано')}
Место: {extracted_data.get('Место совершения деяния', 'не указано')}
Запрашивающее государство: {extracted_data.get('Запрашивающее государство', 'не указано')}
Запрашивающий орган: {extracted_data.get('Запрашивающий орган', 'не указано')}
Наказание: {extracted_data.get('Мера наказания по закону запрашивающего государства', 'не указано')}
Мера пресечения до: {extracted_data.get('Срок меры пресечения', 'не указано')}
"""
    legal_questions = [
        "Каковы безусловные основания для отказа в выдаче (ст. 464 УПК РФ)?",
        "Что такое двойная криминальность и как проверить соответствие деяния по УК РФ?",
        "Сроки давности привлечения к уголовной ответственности за тяжкое преступление (ст. 78 УК РФ).",
        "Каковы требования к гарантиям запрашивающего государства (ст. 462 ч. 3 УПК РФ, ст. 66 Минской конвенции)?",
        "Каков порядок проверки заявлений о политическом преследовании и угрозе пыток (п. 13, 14 Постановления Пленума ВС №11)?",
        "Максимальные сроки содержания под стражей в порядке ст. 109 УПК РФ."
    ]
    rag_results = {}
    for q in legal_questions:
        retrieved = vector_store.search(q) if vector_store else []
        rag_results[q] = "\n\n".join([c for c, _, _ in retrieved[:2]])
    analysis_prompt = f"""
Ты - эксперт по экстрадиционному праву.
Данные по делу:
{context_info}

Расхождения:
{json.dumps(discrepancies, ensure_ascii=False, indent=2)}

Нормы права (RAG):
{json.dumps(rag_results, ensure_ascii=False, indent=2)}

Сформируй разделы:
1. Безусловные основания для отказа (ст. 464 ч. 1 УПК РФ).
2. Факультативные основания (ст. 464 ч. 2 УПК РФ).
3. Условия выдачи (ст. 462 УПК РФ).
4. Проверка меры пресечения (ст. 466 УПК РФ, ст. 109 УПК РФ).
5. Применимый международный договор.
6. Анализ заявлений лица.

Вывод: возможна ли выдача.
"""
    try:
        response = llm_client.chat.completions.create(
            model=SERVED_MODEL_NAME,
            messages=[{"role": "user", "content": analysis_prompt}],
            temperature=0.0, max_tokens=2048,
        )
        analysis_text = response.choices[0].message.content
    except Exception as e:
        analysis_text = f"Ошибка правового анализа: {e}"
    return {"analysis": analysis_text, "rag_used": rag_results}

# ====================== ИТОГОВЫЙ ОТЧЁТ ======================
def generate_full_report(documents_info, all_entities, discrepancies,
                         legal_analysis_result, llm_client) -> str:
    report_date = datetime.now().strftime("%d.%m.%Y %H:%M")
    doc_list_str = "\n".join([f"- {d['filename']} ({d['doc_type']})" for d in documents_info])
    missing_docs = []
    has_guarantees = any(
        d['doc_type'] == "гарантии запрашивающего государства" or
        "гарант" in d['filename'].lower() or
        "гарант" in d.get('text', '').lower()
        for d in documents_info
    )
    if not has_guarantees:
        missing_docs.append("гарантии запрашивающего государства")
    discrepancies_str = ""
    for d in discrepancies:
        discrepancies_str += f"\n**Расхождение:** {d['field']}\n"
        for val in d['values']:
            discrepancies_str += f"  - {val['value']} (документ {val['doc_index']+1})\n"

    report_prompt = f"""
Ты - система формирования финального отчёта по экстрадиционной проверке.
Структура:
- I. Общие сведения.
- II. Полнота комплекта документов.
- III. Извлечённые данные и расхождения.
- IV. Правовой анализ.
- V. Замечания (критические / существенные / рекомендательные).
- VI. Итоговая рекомендация.

Документы:
{doc_list_str}

Отсутствующие: {', '.join(missing_docs) if missing_docs else 'нет'}

Сущности:
{json.dumps(all_entities, ensure_ascii=False, indent=2)}

Расхождения:
{discrepancies_str if discrepancies_str else 'не выявлено'}

Правовой анализ:
{legal_analysis_result.get('analysis', 'не сгенерирован')}
"""
    try:
        response = llm_client.chat.completions.create(
            model=SERVED_MODEL_NAME,
            messages=[{"role": "user", "content": report_prompt}],
            temperature=0.0, max_tokens=8192,
        )
        report = response.choices[0].message.content
    except Exception as e:
        report = f"Ошибка генерации отчёта: {e}"
    return f"Дата формирования отчёта: {report_date}\n\n{report}"

# ====================== ПАЙПЛАЙН АНАЛИЗА ======================
def run_analysis(session_id: str, username: str) -> str:
    with get_db() as conn:
        docs = conn.execute(
            "SELECT filename, file_path, doc_type, extracted_text FROM uploaded_docs WHERE session_id = ?",
            (session_id,)
        ).fetchall()

    if not docs:
        question = "Дай краткий обзор порядка выдачи лиц и оснований отказа (УПК РФ, Минская конвенция)."
        rag_context = ""
        if vector_store is not None:
            try:
                hits = vector_store.search(question, top_k=TOP_K)
                rag_context = "\n\n".join([f"[{src}] {chunk}" for chunk, src, _ in hits])
            except Exception as e:
                print(f"RAG search error: {e}")
        messages = [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": (
                f"Вопрос:\n{question}\n\n"
                + (f"Релевантные нормы:\n{rag_context}\n\n" if rag_context else "")
                + "Ответь по существу, со ссылками на нормы."
            )},
        ]
        try:
            response = client.chat.completions.create(
                model=SERVED_MODEL_NAME, messages=messages,
                temperature=0.2, max_tokens=2048,
            )
            report = response.choices[0].message.content
        except Exception as e:
            report = rag_context or f"LLM недоступен: {e}"
        return f"Дата формирования отчёта: {datetime.now().strftime('%d.%m.%Y %H:%M')}\n\n{report}"

    documents_info, all_entities = [], []
    for doc in docs:
        doc_dict = dict(doc)
        text = doc_dict['extracted_text']
        if not text:
            text = extract_text_from_file(doc_dict['file_path'])
            with get_db() as conn:
                conn.execute(
                    "UPDATE uploaded_docs SET extracted_text = ? WHERE session_id = ? AND filename = ?",
                    (text, session_id, doc_dict['filename'])
                )
        doc_type = doc_dict['doc_type'] or classify_document(text, doc_dict['filename'])
        documents_info.append({
            "filename": doc_dict['filename'],
            "doc_type": doc_type,
            "text": text[:2000],
        })
        entities = extract_entities_from_document(text, doc_type, client)
        all_entities.append(entities)
        with get_db() as conn:
            conn.execute(
                "UPDATE uploaded_docs SET doc_type = ?, entities = ? WHERE session_id = ? AND filename = ?",
                (doc_type, json.dumps(entities, ensure_ascii=False), session_id, doc_dict['filename'])
            )

    discrepancies = cross_check_entities(all_entities)
    aggregated_data = {}
    for ents in all_entities:
        for k, v in ents.items():
            if v and v != "null" and v is not None and k not in aggregated_data:
                aggregated_data[k] = v

    legal_analysis_result = legal_analysis(aggregated_data, discrepancies, client)
    report = generate_full_report(documents_info, all_entities, discrepancies,
                                  legal_analysis_result, client)
    with get_db() as conn:
        conn.execute(
            "INSERT INTO reports (username, session_id, report_text) VALUES (?, ?, ?)",
            (username, session_id, report)
        )
    return report

# ====================== ЧАТ: ВСПОМОГАТЕЛЬНЫЕ ======================
def _ensure_chat(conn, chat_id: Optional[str], username: str,
                 session_id: Optional[str], title_hint: str) -> str:
    if chat_id:
        row = conn.execute(
            "SELECT username, session_id FROM chats WHERE chat_id = ?", (chat_id,)
        ).fetchone()
        if row and row["username"] == username:
            if session_id and not row["session_id"]:
                conn.execute(
                    "UPDATE chats SET session_id = ? WHERE chat_id = ?",
                    (session_id, chat_id),
                )
            return chat_id
    new_id = f"{username}_{uuid.uuid4().hex[:12]}"
    title = (title_hint or "Новый диалог").strip()[:80]
    conn.execute(
        "INSERT INTO chats (chat_id, username, session_id, title) VALUES (?, ?, ?, ?)",
        (new_id, username, session_id, title),
    )
    return new_id

def _save_message(conn, chat_id: str, role: str, content: str):
    conn.execute(
        "INSERT INTO chat_messages (chat_id, role, content) VALUES (?, ?, ?)",
        (chat_id, role, content),
    )
    conn.execute(
        "UPDATE chats SET updated_at = CURRENT_TIMESTAMP WHERE chat_id = ?",
        (chat_id,),
    )

# ====================== МОДЕЛИ ЗАПРОСОВ ======================
class ChatMessage(BaseModel):
    role: str = Field(..., pattern="^(user|assistant|system)$")
    content: str

class AskRequest(BaseModel):
    question: str = Field(..., min_length=1, max_length=8000)
    session_id: Optional[str] = None
    chat_id: Optional[str] = None
    history: Optional[List[ChatMessage]] = None
    use_documents: bool = True
    use_rag: bool = True

class UserRegister(BaseModel):
    username: str
    password: str
    full_name: str | None = None
    organization: str | None = None
    position: str | None = None

class UserLogin(BaseModel):
    username: str
    password: str

class AnalyzeRequest(BaseModel):
    session_id: str

class ChatCreateRequest(BaseModel):
    session_id: Optional[str] = None
    title: Optional[str] = None

# ====================== ЯДРО: СБОРКА СООБЩЕНИЙ ======================
# ====================== ЭВРИСТИКИ ДИАЛОГА ======================
DIALOG_MARKERS = re.compile(
    r"\b(он|она|оно|это|эта|тот|та|там|ещё|еще|его|ему|её|ее|мой|моя|мы|вы)\b",
    re.IGNORECASE,
)

# Мета-вопросы о самом диалоге: «что я спрашивал», «какие данные ты мне дал» и т.п.
META_PATTERNS = re.compile(
    r"("
    r"как(ие|ой|ую)\s+(я|вопрос|данн|информац)|"
    r"что\s+я\s+(у\s+тебя\s+)?(спрашива|успел|говорил|задава|писал)|"
    r"напомни|повтори|"
    r"о\s+ч[её]м\s+мы|"
    r"предыдущ(ий|ие|ая|ее)|"
    r"что\s+ты\s+мне\s+(отправил|дал|писал|прислал)|"
    r"какой\s+(был\s+)?мой\s+предыдущ"
    r")",
    re.IGNORECASE,
)

# Короткие социальные реплики: «ок», «спасибо», «понял» и т.п.
SHORT_SOCIAL = re.compile(
    r"^\s*(ок|окей|ok|okay|угу|ага|да|нет|понял|поняла|ясно|"
    r"спасибо|благодарю|хорошо|принял|принято|ладно|норм|нормально)"
    r"\s*[.!?,]*\s*$",
    re.IGNORECASE,
)


def _is_meta_question(q: str) -> bool:
    """Вопрос о самом диалоге, а не о деле."""
    return bool(META_PATTERNS.search(q or ""))


def _is_social(q: str) -> bool:
    """Короткая социальная реплика без запроса."""
    return bool(SHORT_SOCIAL.match(q or ""))


def _temperature_for(q: str) -> float:
    """
    Для мета/социальных/коротких реплик — выше температура,
    чтобы не было дословных повторов.
    Для правовых вопросов — ниже, чтобы не было фантазий.
    """
    if _is_meta_question(q) or _is_social(q) or len(q or "") < 40:
        return 0.5
    return 0.2

def _load_context_docs(req: AskRequest, username: str) -> str:
    if not (req.use_documents and req.session_id):
        return ""
    with get_db() as conn:
        row = conn.execute(
            "SELECT username FROM upload_sessions WHERE session_id = ?",
            (req.session_id,)
        ).fetchone()
        if not row or row["username"] != username:
            return ""
        docs = conn.execute(
            "SELECT filename, doc_type, extracted_text, file_path "
            "FROM uploaded_docs WHERE session_id = ?",
            (req.session_id,)
        ).fetchall()
        parts = []
        for d in docs:
            text = (d["extracted_text"] or "").strip()
            if not text:
                try:
                    text = extract_text_from_file(d["file_path"])
                    with get_db() as conn2:
                        conn2.execute(
                            "UPDATE uploaded_docs SET extracted_text = ? "
                            "WHERE session_id = ? AND filename = ?",
                            (text, req.session_id, d["filename"])
                        )
                except Exception:
                    text = ""
            parts.append(
                f"### Документ: {d['filename']} "
                f"({d['doc_type'] or 'тип не определён'})\n{text[:5000]}"
            )
        return "\n\n".join(parts)

def _load_rag_context(req: AskRequest) -> Tuple[str, List[str]]:
    if not (req.use_rag and vector_store is not None):
        return "", []
    # Для коротких диалоговых вопросов RAG отключаем — иначе он забивает контекст.
    if len(req.question) < 60 and DIALOG_MARKERS.search(req.question):
        return "", []
    try:
        hits = vector_store.search(req.question, top_k=TOP_K)
        context = "\n\n".join([f"[{src}] {chunk}" for chunk, src, _ in hits])
        sources = [src for _, src, _ in hits]
        return context, sources
    except Exception as e:
        print(f"RAG search error: {e}")
        return "", []

def build_messages_for_llm(
    req: AskRequest,
    username: str,
    context_docs: str,
    rag_context: str,
) -> Tuple[str, List[Dict[str, str]], float]:
    """
    Возвращает (chat_id, messages, temperature).

    Логика:
      1. Определяем/создаём chat_id.
      2. Сохраняем текущий вопрос в БД ДО сборки (чтобы история его видела).
      3. Подгружаем историю из БД (или берём из req.history).
      4. Отрезаем дубль текущего вопроса, если он попал в историю.
      5. В зависимости от типа вопроса:
         - мета-вопрос → только история, без RAG и документов;
         - социальная реплика → короткий ответ + следующий шаг;
         - обычный вопрос → RAG + документы + краткая инструкция.
      6. Возвращаем messages и температуру.
    """
    # ---- 1. chat_id ----
    with get_db() as conn:
        chat_id = _ensure_chat(
            conn, req.chat_id, username, req.session_id, req.question
        )
        # ---- 2. Сохраняем текущий вопрос ДО сборки messages ----
        _save_message(conn, chat_id, "user", req.question)

    # ---- 3. История ----
    if req.history:
        history_msgs = [
            {"role": m.role, "content": m.content}
            for m in req.history[-20:]
        ]
    else:
        with get_db() as conn:
            rows = conn.execute(
                "SELECT role, content FROM chat_messages "
                "WHERE chat_id = ? ORDER BY id ASC LIMIT 40",
                (chat_id,)
            ).fetchall()
            history_msgs = [
                {"role": r["role"], "content": r["content"]}
                for r in rows
            ]

    # ---- 4. Отрезаем дубль текущего вопроса из истории ----
    if history_msgs \
            and history_msgs[-1]["role"] == "user" \
            and history_msgs[-1]["content"] == req.question:
        history_msgs = history_msgs[:-1]

    # ---- 5. Сборка сообщений ----
    messages: List[Dict[str, str]] = [
        {"role": "system", "content": SYSTEM_PROMPT}
    ]
    messages.extend(history_msgs)

    is_meta = _is_meta_question(req.question)
    is_social = _is_social(req.question)

    if is_meta:
        # Мета-вопрос: только история, никаких RAG и документов
        user_parts = [
            f"Мета-вопрос о диалоге:\n{req.question}",
            "Ответь СТРОГО на основе истории диалога выше. "
            "Перечисли вопросы пользователя по порядку, цитируя их дословно. "
            "Если пользователь спрашивает, какие данные ты ему отправлял — "
            "перечисли, что именно было в твоих предыдущих ответах. "
            "НЕ генерируй аналитических справок и отчётов. "
            "Если истории нет — так и скажи.",
        ]
    elif is_social:
        # Социальная реплика: коротко + следующий шаг
        user_parts = [
            f"Реплика пользователя: «{req.question}»",
            "Это короткая социальная реплика (подтверждение/согласие). "
            "Ответь одной фразой и предложи следующий шаг по делу, "
            "которое обсуждается в истории. "
            "НЕ генерируй справок, отчётов и перечней.",
        ]
    else:
        # Обычный содержательный вопрос
        user_parts = [f"Вопрос пользователя:\n{req.question}"]
        if rag_context:
            user_parts.append(
                "Релевантные нормы из правовой базы (RAG):\n" + rag_context
            )
        if context_docs:
            user_parts.append(
                "Контекст загруженных пользователем документов:\n" + context_docs
            )
        user_parts.append(
            "Ответь на вопрос кратко и по делу. Если пользователь явно "
            "просит справку/заключение/таблицу — оформи соответствующим образом."
        )

    messages.append({"role": "user", "content": "\n\n".join(user_parts)})

    temperature = _temperature_for(req.question)
    return chat_id, messages, temperature

# ====================== ЭНДПОИНТЫ ======================
@app.on_event("startup")
def startup():
    init_db()
    start_vllm_server()
    init_rag()
    global client
    client = OpenAI(base_url=f"http://localhost:{VLLM_PORT}/v1", api_key=VLLM_API_KEY)

@app.on_event("shutdown")
def shutdown():
    global vllm_process
    if vllm_process:
        vllm_process.terminate()

# ----- health -----
@app.get("/health")
async def health():
    try:
        with get_db() as conn:
            conn.execute("SELECT 1")
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"DB unavailable: {e}")
    return {"status": "ok"}

# ----- auth -----


@app.post("/register")
@limiter.limit("5/hour")
async def register(request: Request, user: UserRegister):
    username = user.username.strip()
    password = user.password

    if len(username) < 3:
        raise HTTPException(status_code=400, detail="Имя пользователя должно быть не короче 3 символов")
    if len(username) > 50:
        raise HTTPException(status_code=400, detail="Имя пользователя слишком длинное")
    if len(password) < 8:
        raise HTTPException(status_code=400, detail="Пароль должен быть не короче 8 символов")
    if len(password) > 72:
        raise HTTPException(status_code=400, detail="Пароль слишком длинный")
    if get_user(username):
        raise HTTPException(status_code=400, detail="Такое имя пользователя уже занято")

    hashed = get_password_hash(password)
    create_user(username, hashed, user.full_name, user.organization, user.position)
    return {"msg": "User created"}

@app.post("/login")
@limiter.limit("10/minute")
async def login(request: Request, user: UserLogin):
    auth_user = authenticate_user(user.username, user.password)
    if not auth_user:
        raise HTTPException(status_code=401, detail="Invalid credentials")
    token = create_access_token(data={"sub": user.username})
    return {"access_token": token, "token_type": "bearer"}

# ----- upload / analyze / reports -----
@app.post("/upload")
async def upload_documents(
    request: Request,
    files: List[UploadFile] = File(...),
    username: str = Depends(get_current_user),
):
    session_id = f"{username}_{int(time.time())}"
    session_dir = os.path.join(UPLOAD_ROOT, session_id)
    os.makedirs(session_dir, exist_ok=True)
    with get_db() as conn:
        conn.execute(
            "INSERT INTO upload_sessions (username, session_id) VALUES (?, ?)",
            (username, session_id)
        )
    uploaded_files = []
    for file in files:
        file_path = os.path.join(session_dir, file.filename)
        with open(file_path, "wb") as f:
            shutil.copyfileobj(file.file, f)
        uploaded_files.append(file.filename)
        with get_db() as conn:
            conn.execute(
                "INSERT INTO uploaded_docs (session_id, filename, file_path) VALUES (?, ?, ?)",
                (session_id, file.filename, file_path)
            )
    return {"session_id": session_id, "files": uploaded_files}

@app.post("/analyze")
async def analyze(req: AnalyzeRequest, username: str = Depends(get_current_user)):
    session_id = req.session_id
    with get_db() as conn:
        row = conn.execute(
            "SELECT username FROM upload_sessions WHERE session_id = ?", (session_id,)
        ).fetchone()
        if not row or row["username"] != username:
            raise HTTPException(status_code=403, detail="Access denied")
    try:
        report = await asyncio.to_thread(run_analysis, session_id, username)
        return {"report": report}
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/reports")
async def get_reports(username: str = Depends(get_current_user)):
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT id, session_id, created_at,
                   substr(report_text, 1, 200) AS preview,
                   ROW_NUMBER() OVER (PARTITION BY username ORDER BY created_at ASC) AS user_seq
            FROM reports
            WHERE username = ?
            ORDER BY created_at DESC
            """,
            (username,)
        ).fetchall()
    return {"reports": [
        {"id": r["id"], "session_id": r["session_id"],
         "created_at": r["created_at"], "preview": r["preview"],
         "user_seq": r["user_seq"]}
        for r in rows
    ]}

@app.get("/report/{report_id}")
async def get_report(report_id: int, username: str = Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute(
            "SELECT id, session_id, created_at, report_text FROM reports "
            "WHERE id = ? AND username = ?",
            (report_id, username)
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Report not found")
        return {"id": row["id"], "session_id": row["session_id"],
                "created_at": row["created_at"], "report": row["report_text"]}

@app.delete("/report/{report_id}")
async def delete_report(report_id: int, username: str = Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute(
            "SELECT username FROM reports WHERE id = ? AND username = ?",
            (report_id, username)
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Report not found")
        conn.execute("DELETE FROM reports WHERE id = ?", (report_id,))
    return {"ok": True}

# ----- chats -----
@app.get("/chats")
async def list_chats(username: str = Depends(get_current_user)):
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT c.chat_id, c.session_id, c.title, c.created_at, c.updated_at,
                   (SELECT COUNT(*) FROM chat_messages m WHERE m.chat_id = c.chat_id) AS messages_count
            FROM chats c
            WHERE c.username = ?
            ORDER BY c.updated_at DESC
            """,
            (username,)
        ).fetchall()
    return {"chats": [dict(r) for r in rows]}

@app.get("/chats/{chat_id}/messages")
async def get_chat_messages(chat_id: str, username: str = Depends(get_current_user)):
    with get_db() as conn:
        owner = conn.execute(
            "SELECT username FROM chats WHERE chat_id = ?", (chat_id,)
        ).fetchone()
        if not owner or owner["username"] != username:
            raise HTTPException(status_code=404, detail="Chat not found")
        rows = conn.execute(
            "SELECT role, content, created_at FROM chat_messages "
            "WHERE chat_id = ? ORDER BY id ASC",
            (chat_id,)
        ).fetchall()
    return {"messages": [dict(r) for r in rows]}

@app.post("/chats")
async def create_chat(req: ChatCreateRequest, username: str = Depends(get_current_user)):
    chat_id = f"{username}_{uuid.uuid4().hex[:12]}"
    with get_db() as conn:
        conn.execute(
            "INSERT INTO chats (chat_id, username, session_id, title) VALUES (?, ?, ?, ?)",
            (chat_id, username, req.session_id, req.title or "Новый диалог")
        )
    return {"chat_id": chat_id}

@app.delete("/chats/{chat_id}")
async def delete_chat(chat_id: str, username: str = Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute(
            "SELECT username FROM chats WHERE chat_id = ?", (chat_id,)
        ).fetchone()
        if not row or row["username"] != username:
            raise HTTPException(status_code=404, detail="Chat not found")
        conn.execute("DELETE FROM chat_messages WHERE chat_id = ?", (chat_id,))
        conn.execute("DELETE FROM chats WHERE chat_id = ?", (chat_id,))
    return {"ok": True}

# ----- stats -----
@app.get("/stats")
async def get_stats(username: str = Depends(get_current_user)):
    with get_db() as conn:
        total_reports = conn.execute(
            "SELECT COUNT(*) AS c FROM reports WHERE username = ?", (username,)
        ).fetchone()["c"]
        week_ago = (datetime.now() - timedelta(days=7)).strftime("%Y-%m-%d %H:%M:%S")
        reports_week = conn.execute(
            "SELECT COUNT(*) AS c FROM reports WHERE username = ? AND created_at >= ?",
            (username, week_ago)
        ).fetchone()["c"]
        total_docs = conn.execute(
            "SELECT COUNT(*) AS c FROM uploaded_docs d "
            "JOIN upload_sessions s ON s.session_id = d.session_id "
            "WHERE s.username = ?",
            (username,)
        ).fetchone()["c"]
        total_sessions = conn.execute(
            "SELECT COUNT(*) AS c FROM upload_sessions WHERE username = ?", (username,)
        ).fetchone()["c"]
    return {
        "total_reports": total_reports,
        "reports_this_week": reports_week,
        "total_documents": total_docs,
        "total_sessions": total_sessions,
    }

# ----- ask -----
@app.post("/ask")
async def ask(req: AskRequest, username: str = Depends(get_current_user)):
    context_docs = _load_context_docs(req, username)
    rag_context, rag_sources = _load_rag_context(req)

    chat_id, messages, temperature = build_messages_for_llm(
        req, username, context_docs, rag_context
    )
    # ВАЖНО: _save_message(conn, chat_id, "user", req.question) здесь больше НЕ нужен —
    # он уже вызван внутри build_messages_for_llm.

    try:
        response = await asyncio.to_thread(
            client.chat.completions.create,
            model=SERVED_MODEL_NAME,
            messages=messages,
            temperature=temperature,       # <-- теперь из функции
            max_tokens=4096,
        )
        answer = response.choices[0].message.content
    except Exception as e:
        import traceback
        traceback.print_exc()
        if rag_context:
            answer = "LLM недоступен, привожу релевантные нормы:\n\n" + rag_context
        else:
            raise HTTPException(status_code=500, detail=f"LLM error: {e}")

    with get_db() as conn:
        _save_message(conn, chat_id, "assistant", answer)

    return {
        "answer": answer,
        "chat_id": chat_id,
        "used_rag": bool(rag_context),
        "used_documents": bool(context_docs),
        "rag_sources": rag_sources,
    }
# ----- ask_stream -----

@app.post("/ask_stream")
async def ask_stream(req: AskRequest, username: str = Depends(get_current_user)):
    context_docs = _load_context_docs(req, username)
    rag_context, _ = _load_rag_context(req)

    chat_id, messages, temperature = build_messages_for_llm(
        req, username, context_docs, rag_context
    )

    def event_generator():
        yield f"__CHAT_ID__:{chat_id}\n\n"
        full_answer = ""
        try:
            stream = client.chat.completions.create(
                model=SERVED_MODEL_NAME,
                messages=messages,
                temperature=temperature,   # <-- теперь из функции
                max_tokens=8192,
                stream=True,
            )
            for chunk in stream:
                try:
                    delta = chunk.choices[0].delta.content or ""
                except Exception:
                    delta = ""
                if delta:
                    full_answer += delta
                    yield delta
        except Exception as e:
            yield f"\n[Ошибка LLM: {e}]"
        finally:
            with get_db() as conn2:
                _save_message(conn2, chat_id, "assistant", full_answer)

    return StreamingResponse(event_generator(), media_type="text/plain; charset=utf-8")

if __name__ == "__main__":
    uvicorn.run(app, host="127.0.0.1", port=8000)